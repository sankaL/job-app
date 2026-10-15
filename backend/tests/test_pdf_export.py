from __future__ import annotations

import io
import re

import pytest

from app.services import pdf_export


def _personal_info() -> dict[str, str]:
    return {
        "name": "Alex Example",
        "email": "alex@example.com",
        "phone": "555-0100",
        "address": "Toronto, ON",
        "linkedin_url": "https://linkedin.com/in/alex-example",
    }


def test_normalize_markdown_replaces_legacy_placeholder_header():
    normalized = pdf_export._normalize_markdown_for_export(
        "# (Name)\ninvite-only@example.com\n\n## Summary\nBuilt backend systems.\n",
        _personal_info(),
    )

    assert normalized.startswith(
        "# Alex Example\nalex@example.com | 555-0100 | Toronto, ON | linkedin.com/in/alex-example\n\n## Summary\n"
    )
    assert "(Name)" not in normalized
    assert "invite-only@example.com" not in normalized


def test_build_html_renders_one_header_for_existing_assembly_header():
    markdown_content = (
        "# Alex Example\n"
        "alex@example.com | 555-0100 | Toronto, ON | linkedin.com/in/alex-example\n\n"
        "## Summary\n"
        "Built backend systems.\n"
    )

    html = pdf_export._build_html(markdown_content, pdf_export.LAYOUT_PRESETS[0])

    assert html.count("<header class=\"resume-header\">") == 0
    assert html.count("<header class='resume-header'>") == 1
    assert html.count("<h1>Alex Example</h1>") == 1
    assert "## Summary" not in html
    assert "Built backend\u00a0systems." in html


def test_normalize_markdown_replaces_plain_text_header_matching_profile():
    normalized = pdf_export._normalize_markdown_for_export(
        "Alex Example\nalex@example.com | 555-0100\n\n## Summary\nBuilt backend systems.\n",
        _personal_info(),
    )

    assert normalized.startswith(
        "# Alex Example\nalex@example.com | 555-0100 | Toronto, ON | linkedin.com/in/alex-example\n\n## Summary\n"
    )
    assert "\n\nAlex Example\nalex@example.com | 555-0100\n" not in normalized


def test_build_html_uses_point_spacing_units():
    preset = pdf_export.LAYOUT_PRESETS[0]
    html = pdf_export._build_html(
        "## Summary\nBuilt backend systems.\n",
        preset,
    )
    major_section_gap = max(
        preset.body_font_size * 0.6 * preset.section_spacing_scale,
        preset.section_margin_top * 1.1,
    )
    header_gap = max(major_section_gap, preset.contact_to_first_section_margin * 1.1)

    assert f"margin-bottom: {header_gap:.2f}pt;" in html
    assert f"margin-top: {major_section_gap:.2f}pt;" in html
    assert f"margin: 0 0 {preset.section_header_content_gap * 1.1:.2f}pt 0;" in html
    assert f"margin: 0 0 {preset.paragraph_margin * 1.08:.2f}pt 0;" in html
    assert f"margin: 0 0 {preset.paragraph_margin * 1.08:.2f}pt {preset.bullet_indent}pt;" in html
    assert "margin-top: 0;" in html
    assert 'data-density="sparse"' in html
    assert f"{preset.section_margin_top}rem" not in html
    assert f"{preset.section_header_content_gap}rem" not in html


def test_render_content_blocks_unwraps_single_item_list_markup_inside_bullets():
    html = pdf_export._render_content_blocks(["- * nested bullet"])

    assert html == "<ul><li>nested\u00a0bullet</li></ul>"


def test_render_content_blocks_preserves_literal_star_content_inside_bullets():
    html = pdf_export._render_content_blocks(["- *nix operations"])

    assert "*nix\u00a0operations" in html
    assert "<ul><li><ul>" not in html


def test_calculate_content_density_metrics_flags_dense_and_sparse_documents():
    dense_metrics = pdf_export._calculate_content_density_metrics(
        (
            "## Experience\n"
            "- one\n- two\n- three\n- four\n- five\n"
            "## Projects\n"
            "- six\n- seven\n- eight\n- nine\n- ten\n"
            "## Skills\n"
            "- python\n- sql\n- aws\n- docker\n- linux\n"
        )
    )
    sparse_metrics = pdf_export._calculate_content_density_metrics(
        "## Summary\nBuilt backend systems.\n## Skills\nPython | SQL\n"
    )

    assert dense_metrics["is_dense"] is True
    assert dense_metrics["density_label"] == "dense"
    assert sparse_metrics["is_sparse"] is True
    assert sparse_metrics["density_label"] == "sparse"


def test_layout_presets_keep_readable_minimums():
    assert all(preset.body_font_size >= pdf_export.MIN_READABLE_BODY_FONT_SIZE for preset in pdf_export.LAYOUT_PRESETS)
    assert all(preset.line_height >= pdf_export.MIN_READABLE_LINE_HEIGHT for preset in pdf_export.LAYOUT_PRESETS)


def test_generate_pdf_autofit_retries_until_target_page_count_is_met(monkeypatch):
    seen_presets: list[int] = []

    def fake_render_html_to_pdf(html_content: str) -> tuple[bytes, int]:
        preset_index = int(re.search(r'data-preset="(\d+)"', html_content).group(1))
        seen_presets.append(preset_index)
        page_counts = {0: 3, 1: 2, 2: 1}
        return (f"preset-{preset_index}".encode(), page_counts.get(preset_index, 1))

    monkeypatch.setattr(pdf_export, "_render_html_to_pdf", fake_render_html_to_pdf)

    pdf_bytes = pdf_export._generate_pdf_with_autofit_sync(
        "## Summary\nBuilt backend systems.\n",
        _personal_info(),
        "1_page",
    )

    assert pdf_bytes == b"preset-2"
    assert seen_presets[:3] == [0, 1, 2]
    assert all(index == 2 for index in seen_presets[3:])


def test_generate_pdf_autofit_prefers_tighter_density_before_smaller_font(monkeypatch):
    seen_presets: list[int] = []

    def fake_render_html_to_pdf(html_content: str) -> tuple[bytes, int]:
        preset_index = int(re.search(r'data-preset="(\d+)"', html_content).group(1))
        seen_presets.append(preset_index)
        page_counts = {0: 2, 1: 1, 2: 1}
        return (f"preset-{preset_index}".encode(), page_counts.get(preset_index, 1))

    monkeypatch.setattr(pdf_export, "_render_html_to_pdf", fake_render_html_to_pdf)

    pdf_bytes = pdf_export._generate_pdf_with_autofit_sync(
        "## Summary\nBuilt backend systems.\n",
        _personal_info(),
        "1_page",
    )

    assert pdf_bytes == b"preset-1"
    assert seen_presets[:2] == [0, 1]
    assert all(index == 1 for index in seen_presets[2:])


def test_generate_pdf_autofit_stops_at_minimum_preset(monkeypatch):
    seen_presets: list[int] = []

    def fake_render_html_to_pdf(html_content: str) -> tuple[bytes, int]:
        preset_index = int(re.search(r'data-preset="(\d+)"', html_content).group(1))
        seen_presets.append(preset_index)
        return (f"preset-{preset_index}".encode(), 5)

    monkeypatch.setattr(pdf_export, "_render_html_to_pdf", fake_render_html_to_pdf)

    pdf_bytes = pdf_export._generate_pdf_with_autofit_sync(
        "## Summary\nBuilt backend systems.\n",
        _personal_info(),
        "1_page",
    )

    min_preset_index = len(pdf_export.LAYOUT_PRESETS) - 1
    assert pdf_bytes == f"preset-{min_preset_index}".encode()
    assert seen_presets == list(range(len(pdf_export.LAYOUT_PRESETS)))


def test_generate_pdf_one_page_validation_tries_roomier_variant_before_export(monkeypatch):
    seen_fonts: list[float] = []

    monkeypatch.setattr(
        pdf_export,
        "LAYOUT_PRESETS",
        [pdf_export.LayoutPreset(body_font_size=10.0, line_height=1.10, page_margin=0.3, spacing_scale=0.7)],
    )

    def fake_render_html_to_pdf(html_content: str) -> tuple[bytes, int]:
        body_font_size = float(re.search(r"font-size:\s*([0-9.]+)pt;", html_content).group(1))
        seen_fonts.append(body_font_size)
        page_count = 1 if body_font_size <= 10.22 else 2
        return (f"font-{body_font_size:.2f}".encode(), page_count)

    monkeypatch.setattr(pdf_export, "_render_html_to_pdf", fake_render_html_to_pdf)

    pdf_bytes = pdf_export._generate_pdf_with_autofit_sync(
        "## Summary\nBuilt backend systems.\n",
        _personal_info(),
        "1_page",
    )

    assert pdf_bytes == b"font-10.22"
    assert seen_fonts[0] == 10.0
    assert 10.22 in seen_fonts
    assert 10.44 in seen_fonts
    assert len(seen_fonts) >= 3


def test_generate_pdf_one_page_validation_adds_section_spacing_when_room_exists(monkeypatch):
    seen_section_tops: list[float] = []
    accepted_section_tops: list[float] = []

    monkeypatch.setattr(
        pdf_export,
        "LAYOUT_PRESETS",
        [pdf_export.LayoutPreset(body_font_size=10.0, line_height=1.10, page_margin=0.3, spacing_scale=0.7)],
    )

    def fake_render_html_to_pdf(html_content: str) -> tuple[bytes, int]:
        section_margin_top = float(re.search(r"\.resume-section \{\s*margin-top: ([0-9.]+)pt;", html_content).group(1))
        body_font_size = float(re.search(r"font-size:\s*([0-9.]+)pt;", html_content).group(1))
        seen_section_tops.append(section_margin_top)
        # Allow section spacing growth, but force any font growth to overflow.
        page_count = 1 if body_font_size <= 10.0 else 2
        if page_count == 1:
            accepted_section_tops.append(section_margin_top)
        return (f"section-top-{section_margin_top:.2f}".encode(), page_count)

    monkeypatch.setattr(pdf_export, "_render_html_to_pdf", fake_render_html_to_pdf)

    pdf_bytes = pdf_export._generate_pdf_with_autofit_sync(
        "## Summary\nBuilt backend systems.\n",
        _personal_info(),
        "1_page",
    )

    assert seen_section_tops[0] > 0
    assert max(accepted_section_tops) > seen_section_tops[0]
    assert pdf_bytes == f"section-top-{max(accepted_section_tops):.2f}".encode()


def test_build_html_marks_dense_documents_in_html():
    html = pdf_export._build_html(
        (
            "## Experience\n"
            "- one\n- two\n- three\n- four\n- five\n"
            "## Projects\n"
            "- six\n- seven\n- eight\n- nine\n- ten\n"
            "## Skills\n"
            "- python\n- sql\n- aws\n- docker\n- linux\n"
        ),
        pdf_export.LAYOUT_PRESETS[0],
    )

    assert 'data-density="dense"' in html


def test_build_html_keeps_structured_entry_case_and_font_hierarchy():
    html = pdf_export._build_html(
        (
            "## Professional Experience\n"
            "Deloitte Canada | Toronto, Ontario, Canada\n"
            "Manager, Product Engineering | Jan 2022 - Present\n"
            "- Led platform work.\n"
        ),
        pdf_export.LAYOUT_PRESETS[0],
    )

    section_heading_size = pdf_export.LAYOUT_PRESETS[0].section_heading_size
    expected_primary_size = round(
        min(
            section_heading_size - 0.18,
            max(
                round(pdf_export.LAYOUT_PRESETS[0].body_font_size * 0.86, 2) + 0.32,
                pdf_export.LAYOUT_PRESETS[0].body_font_size * 0.90,
            ),
        ),
        2,
    )
    expected_content_size = round(pdf_export.LAYOUT_PRESETS[0].body_font_size * 0.86, 2)

    assert "DELOITTE CANADA" not in html
    assert "Deloitte Canada" in html
    assert expected_primary_size <= section_heading_size
    assert f"font-size: {expected_primary_size}pt;" in html
    assert f"font-size: {expected_content_size}pt;" in html


def test_build_html_bolds_only_professional_experience_role_title_split_rows():
    html = pdf_export._build_html(
        (
            "## Professional Experience\n"
            "Acme Corp | Toronto, ON\n"
            "Senior Data Architect | Jan 2020 - Present\n"
        ),
        pdf_export.LAYOUT_PRESETS[0],
    )

    assert "<span class='structured-left structured-left-primary'>Acme Corp</span>" in html
    assert "<span class='structured-left structured-left-secondary'>Senior Data Architect</span>" in html
    assert "Senior Data Architect" in html
    assert "Acme Corp" in html


def test_build_html_does_not_bold_date_split_rows_outside_professional_experience():
    html = pdf_export._build_html(
        (
            "## Summary\n"
            "Program Timeline | Jan 2020 - Present\n"
        ),
        pdf_export.LAYOUT_PRESETS[0],
    )

    assert "structured-entry structured-entry-summary" not in html


def test_build_html_preserves_blank_line_separation_between_split_row_groups():
    html = pdf_export._build_html(
        (
            "## Professional Experience\n"
            "Acme Corp | Toronto, ON\n"
            "Senior Data Architect | Jan 2020 - Present\n\n"
            "Beta Corp | Ottawa, ON\n"
            "Lead Engineer | Jan 2018 - Dec 2019\n"
        ),
        pdf_export.LAYOUT_PRESETS[0],
    )

    assert html.count("class='structured-entry structured-entry-professional_experience'") == 2


def test_build_export_document_reuses_one_normalized_header():
    document = pdf_export._build_export_document(
        "Alex Example\nalex@example.com | 555-0100\n\n## Summary\nBuilt backend systems.\n",
        _personal_info(),
    )

    assert document.header is not None
    assert document.header.name == "Alex Example"
    assert document.header.contact_line == "alex@example.com | 555-0100 | Toronto, ON | linkedin.com/in/alex-example"
    assert len(document.sections) == 1
    assert document.sections[0].heading == "Summary"


def test_render_docx_sync_uses_letter_page_size_and_expected_margins():
    from docx import Document

    docx_bytes = pdf_export._render_docx_sync(
        "## Summary\nBuilt backend systems.\n",
        _personal_info(),
        "2_page",
    )
    document = Document(io.BytesIO(docx_bytes))
    section = document.sections[0]
    layout = pdf_export.DOCX_LAYOUT_PRESETS["2_page"]

    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert round(section.top_margin.inches, 2) == round(layout.page_margin, 2)
    assert round(section.left_margin.inches, 2) == round(layout.page_margin, 2)


def test_render_docx_sync_adds_clear_header_and_major_section_spacing():
    from docx import Document

    docx_bytes = pdf_export._render_docx_sync(
        (
            "# Alex Example\n"
            "alex@example.com | 555-0100 | Toronto, ON\n\n"
            "## Summary\n"
            "Built backend systems.\n\n"
            "## Skills\n"
            "- Python\n"
        ),
        _personal_info(),
        "2_page",
    )
    document = Document(io.BytesIO(docx_bytes))
    layout = pdf_export.DOCX_LAYOUT_PRESETS["2_page"]
    header_section_gap = max(layout.body_font_size * layout.line_spacing * 0.95, layout.header_spacing_after)
    major_section_gap = max(layout.body_font_size * layout.line_spacing * 0.95, layout.section_spacing_before)
    contact_paragraph = next(paragraph for paragraph in document.paragraphs if "alex@example.com" in paragraph.text)
    skills_heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "SKILLS")

    assert round(contact_paragraph.paragraph_format.space_after.pt, 2) >= round(header_section_gap, 2)
    assert round(skills_heading.paragraph_format.space_before.pt, 2) >= round(major_section_gap, 2)


def test_render_docx_sync_renders_header_bullets_and_split_rows_without_tables():
    from docx import Document

    docx_bytes = pdf_export._render_docx_sync(
        (
            "# Alex Example\n"
            "alex@example.com | 555-0100 | Toronto, ON\n\n"
            "## Professional Experience\n"
            "Acme Corp | Toronto, ON\n"
            "Senior Data Architect | Jan 2020 - Present\n"
            "- Led migration program\n"
        ),
        _personal_info(),
        "1_page",
    )
    document = Document(io.BytesIO(docx_bytes))
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]

    assert document.tables == []
    assert paragraph_text[0] == "Alex Example"
    assert any("PROFESSIONAL EXPERIENCE" == text for text in paragraph_text)
    assert any("Acme Corp\tToronto, ON" == text for text in paragraph_text)
    assert any("Senior Data Architect\tJan 2020 - Present" == text for text in paragraph_text)
    assert any("Led migration\u00a0program" in text for text in paragraph_text)


def test_render_docx_sync_bolds_only_professional_experience_role_title_split_rows():
    from docx import Document

    docx_bytes = pdf_export._render_docx_sync(
        (
            "## Professional Experience\n"
            "Acme Corp | Toronto, ON\n"
            "Senior Data Architect | Jan 2020 - Present\n"
            "## Summary\n"
            "Portfolio Lead | 2024 - Present\n"
        ),
        _personal_info(),
        "1_page",
    )
    document = Document(io.BytesIO(docx_bytes))
    company_paragraph = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Acme Corp\tToronto, ON")
    role_paragraph = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Senior Data Architect\tJan 2020 - Present")
    summary_paragraph = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Portfolio Lead\t2024 - Present")

    assert company_paragraph.runs[0].bold is True
    assert role_paragraph.runs[0].italic is True
    assert role_paragraph.runs[0].bold is True
    assert summary_paragraph.runs[0].bold in {False, None}


def test_resolve_docx_layout_adjusts_for_dense_and_sparse_documents():
    dense = pdf_export._resolve_docx_layout("1_page", {"is_dense": True, "is_sparse": False})
    sparse = pdf_export._resolve_docx_layout("1_page", {"is_dense": False, "is_sparse": True})
    baseline = pdf_export.DOCX_LAYOUT_PRESETS["1_page"]

    assert dense.body_font_size < baseline.body_font_size
    assert dense.paragraph_spacing < baseline.paragraph_spacing
    assert sparse.body_font_size > baseline.body_font_size
    assert sparse.section_spacing_before > baseline.section_spacing_before


_STYLE_SAMPLE_MARKDOWN = (
    "## Summary\nBuilt backend systems.\n\n"
    "## Professional Experience\n"
    "Acme Corp | Toronto, ON\n"
    "Senior Data Architect | Jan 2020 - Present\n"
    "- Led platform work.\n"
)


def test_classic_style_html_matches_unstyled_output():
    from app.services.export_styles import resolve_style

    preset = pdf_export.LAYOUT_PRESETS[0]
    unstyled = pdf_export._build_html(_STYLE_SAMPLE_MARKDOWN, preset)
    classic = pdf_export._build_html(_STYLE_SAMPLE_MARKDOWN, preset, style=resolve_style("classic"))

    assert classic == unstyled
    # The image has no Georgia, so Classic bundles a Georgia-metric serif instead of falling back to DejaVu.
    assert "font-family: 'Resume Gelasio', 'Georgia', 'Times New Roman', serif;" in classic
    assert "gelasio-400-normal.ttf" in classic
    for face in resolve_style("classic").font_faces:
        assert face.path.is_file(), face.path


def test_non_classic_styles_override_fonts_with_bundled_files():
    from app.services.export_styles import list_export_styles

    preset = pdf_export.LAYOUT_PRESETS[0]
    for style in list_export_styles():
        if style.is_classic:
            continue
        html = pdf_export._build_html(_STYLE_SAMPLE_MARKDOWN, preset, style=style)
        assert f'data-style="{style.key}"' in html
        assert "@font-face" in html
        assert style.font_faces, style.key
        for face in style.font_faces:
            assert face.path.is_file(), face.path


def test_unknown_export_style_fails_closed():
    from app.services.export_styles import resolve_style

    with pytest.raises(ValueError, match="Unknown export style"):
        resolve_style("fancy")
    assert resolve_style(None).key == "classic"
    assert resolve_style(" Modern ").key == "modern"


def test_every_style_renders_pdf_and_docx():
    from docx import Document

    from app.services.export_styles import list_export_styles

    for style in list_export_styles():
        pdf_bytes = pdf_export._generate_pdf_with_autofit_sync(
            _STYLE_SAMPLE_MARKDOWN, _personal_info(), "1_page", style
        )
        assert pdf_bytes.startswith(b"%PDF-"), style.key

        docx_bytes = pdf_export._render_docx_sync(_STYLE_SAMPLE_MARKDOWN, _personal_info(), "1_page", style)
        document = Document(io.BytesIO(docx_bytes))
        assert document.styles["Normal"].font.name == style.docx_font
        headings = [p.text for p in document.paragraphs if p.text in {"SUMMARY", "Summary"}]
        expected = "SUMMARY" if style.heading_transform == "uppercase" else "Summary"
        assert headings == [expected], style.key


_EXPERIENCE_ALIAS_MARKDOWN = (
    "## Experience\n"
    "Acme Corp | Toronto, ON\n"
    "Senior Data Architect | Jan 2020 - Present\n"
    "- Led platform work.\n\n"
    "## Projects\n"
    "- Built Gust, a voice-first task manager that turns speech into tasks.\n"
    "- Ledger: a budgeting tool for small teams.\n\n"
    "## Skills\n"
    "- AI evaluation: LLM-as-judge, RAG evaluation\n"
    "- Python, SQL\n"
)


def test_experience_heading_alias_uses_structured_entry_layout():
    html = pdf_export._build_html(_EXPERIENCE_ALIAS_MARKDOWN, pdf_export.LAYOUT_PRESETS[0])

    assert html.count("class='structured-entry structured-entry-professional_experience'") == 1
    assert ".structured-entry-professional_experience .structured-left-secondary {\n      font-weight: 700;" in html


def test_experience_heading_alias_falls_back_to_markdown_when_entries_do_not_fit():
    document = pdf_export._build_export_document("## Experience\nWorked across many teams.\n")

    assert document.sections[0].heading == "Experience"
    assert not any(isinstance(block, pdf_export.ExportStructuredEntry) for block in document.sections[0].blocks)


def test_project_names_lead_and_skill_labels_are_bold():
    html = pdf_export._build_html(_EXPERIENCE_ALIAS_MARKDOWN, pdf_export.LAYOUT_PRESETS[0])

    assert "<strong>Gust:</strong> Built a voice-first task manager" in html
    assert "<strong>Ledger:</strong> a budgeting tool" in html
    assert "<strong>AI evaluation:</strong> LLM-as-judge" in html
    assert "<strong>Python" not in html


def test_project_name_lead_keeps_every_word():
    text = "Built AbleRoute, a layered extraction pipeline for delivery photos."

    assert pdf_export._lead_with_project_name(text) == "**AbleRoute:** Built a layered extraction pipeline for delivery photos."
    assert pdf_export._lead_with_project_name("Built the RAG suite for retrieval.") == "Built the RAG suite for retrieval."


def test_skills_paragraph_lines_keep_bold_labels_on_separate_lines():
    document = pdf_export._build_export_document("## Skills\nLanguages: Python, SQL\nCloud: AWS, Azure\n")
    paragraph = document.sections[0].blocks[0]

    assert "<strong>Languages:</strong> Python, SQL<br />" in paragraph.html_fragment
    assert "<strong>Cloud:</strong> AWS," in paragraph.html_fragment


def test_bullets_and_paragraphs_never_end_on_one_short_word():
    nbsp = "\u00a0"

    assert pdf_export._bind_trailing_words("shipping every release on budget.") == f"shipping every release{nbsp}on{nbsp}budget."
    assert pdf_export._bind_trailing_words("cutting work from days to hours.") == f"cutting work from{nbsp}days{nbsp}to{nbsp}hours."
    assert pdf_export._bind_trailing_words("experiment management.") == f"experiment{nbsp}management."
    assert pdf_export._bind_trailing_words("<strong>Gust:</strong> a tool") == f"<strong>Gust:</strong> a{nbsp}tool"
    assert pdf_export._bind_trailing_words("short") == "short"


def test_linkedin_contact_keeps_the_domain():
    assert pdf_export._format_linkedin_value("https://www.linkedin.com/in/sankal/") == "linkedin.com/in/sankal"
    assert pdf_export._format_linkedin_value("ca.linkedin.com/in/sankal") == "linkedin.com/in/sankal"


def test_render_docx_sync_bolds_project_names_and_skill_labels():
    from docx import Document

    docx_bytes = pdf_export._render_docx_sync(_EXPERIENCE_ALIAS_MARKDOWN, _personal_info(), "1_page")
    document = Document(io.BytesIO(docx_bytes))
    project = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Gust:"))
    skill = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("AI evaluation:"))
    role = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Senior Data Architect"))

    assert project.runs[0].text == "Gust:" and project.runs[0].bold is True
    assert skill.runs[0].text == "AI evaluation:" and skill.runs[0].bold is True
    assert role.runs[0].bold is True
